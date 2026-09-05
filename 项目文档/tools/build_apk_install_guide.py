from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"D:\pico 4\项目文档\翼揽无余_APK安装与运行指南.docx")

NAVY = "17365D"
BLUE = "2E74B5"
CYAN = "18A7B5"
LIGHT_BLUE = "E8F1F8"
LIGHT_CYAN = "E8F7F8"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
DARK = "1D2939"
WHITE = "FFFFFF"
GOLD = "B7791F"
LIGHT_GOLD = "FFF6E0"
RED = "B42318"
LIGHT_RED = "FEECEB"
GREEN = "16794B"
LIGHT_GREEN = "EAF7F0"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(run, size=None, bold=None, color=None, italic=None, font="Calibri"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = rgb(color)
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D0D5DD", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    total = sum(widths_dxa)
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for i, cell in enumerate(row.cells):
            width = widths_dxa[min(i, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def format_cell_text(cell, size=9.5, color=DARK, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.15
        for run in paragraph.runs:
            set_run_font(run, size=size, color=color, bold=bold)


def add_table(doc, headers, rows, widths_dxa, header_fill=NAVY, first_col_bold=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        set_cell_shading(cell, header_fill)
        format_cell_text(cell, size=9.5, color=WHITE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cells[i].text = str(value)
            if first_col_bold and i == 0:
                set_cell_shading(cells[i], LIGHT_GRAY)
            format_cell_text(cells[i], bold=(first_col_bold and i == 0))
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_callout(doc, label, text, fill=LIGHT_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(f"{label}  ")
    set_run_font(r, size=10.5, bold=True, color=accent)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=DARK)
    set_table_geometry(table, [9360])
    set_table_borders(table, color=accent, size="8")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_body(doc, text, bold_prefix=None, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=10.5, bold=True, color=DARK)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=10.5, color=DARK)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.5, color=DARK)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.375 if level == 0 else 0.65)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    p.add_run(text)
    for run in p.runs:
        set_run_font(run, size=10.5, color=DARK)
    return p


def add_numbered_sequence(doc, items):
    """Create a genuine Word numbering sequence that restarts at 1."""
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    abstract_id = (max(abstract_ids) + 1) if abstract_ids else 0
    num_id = (max(num_ids) + 1) if num_ids else 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "right")
    lvl.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "600")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "600")
    ind.set(qn("w:hanging"), "300")
    p_pr.append(ind)
    lvl.append(p_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_ref = OxmlElement("w:abstractNumId")
    abs_ref.set(qn("w:val"), str(abstract_id))
    num.append(abs_ref)
    numbering.append(num)

    for text in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.2
        p_pr = p._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_ref = OxmlElement("w:numId")
        num_ref.set(qn("w:val"), str(num_id))
        num_pr.append(ilvl)
        num_pr.append(num_ref)
        p_pr.append(num_pr)
        run = p.add_run(text)
        set_run_font(run, size=10.5, color=DARK)


def add_command(doc, command):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F7F8FA")
    set_cell_margins(cell, top=110, start=180, bottom=110, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(command)
    set_run_font(run, size=9.5, color=NAVY, font="Consolas")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    set_table_geometry(table, [9360])
    set_table_borders(table, color="D0D5DD")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = paragraph.add_run("第 ")
    set_run_font(r, size=9, color=MID_GRAY)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r._r.append(fld_begin)
    r._r.append(instr)
    r._r.append(fld_end)
    r2 = paragraph.add_run(" 页")
    set_run_font(r2, size=9, color=MID_GRAY)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.82)
section.bottom_margin = Inches(0.78)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.header_distance = Inches(0.4)
section.footer_distance = Inches(0.4)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for style_name, size, color, before, after in (
    ("Heading 1", 16, NAVY, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 11.5, NAVY, 9, 4),
):
    style = styles[style_name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = rgb(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

# Running header and footer
header_p = section.header.paragraphs[0]
header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
header_p.paragraph_format.space_after = Pt(0)
r = header_p.add_run("翼揽无余  |  APK安装与运行指南")
set_run_font(r, size=9, bold=True, color=MID_GRAY)

footer_table = section.footer.add_table(rows=1, cols=2, width=Inches(6.5))
footer_table.alignment = WD_TABLE_ALIGNMENT.CENTER
footer_table.autofit = False
left = footer_table.cell(0, 0)
right = footer_table.cell(0, 1)
left.text = "比赛提交材料  ·  文档版本 V1.0"
format_cell_text(left, size=8.5, color=MID_GRAY)
add_page_number(right.paragraphs[0])
set_table_geometry(footer_table, [6500, 2860], indent_dxa=0)
for cell in (left, right):
    set_cell_margins(cell, top=0, start=0, bottom=0, end=0)

# Title block
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(34)
p.paragraph_format.space_after = Pt(4)
r = p.add_run("比赛作品安装说明")
set_run_font(r, size=11, bold=True, color=CYAN)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(6)
r = p.add_run("《翼揽无余》APK安装与运行指南")
set_run_font(r, size=26, bold=True, color=NAVY)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(18)
r = p.add_run("适用于 PICO 4 Ultra 的航空科普沉浸式交互应用")
set_run_font(r, size=13, color=MID_GRAY)

add_callout(
    doc,
    "指定验收设备",
    "PICO 4 Ultra 系列，建议系统版本 5.14.0 或以上。当前比赛APK包含PICO视频透视和手势交互能力，不建议使用普通安卓设备或其他品牌头显进行验收。",
    fill=LIGHT_CYAN,
    accent=CYAN,
)

doc.add_heading("一、提交文件信息", level=1)
add_table(
    doc,
    ["项目", "当前APK信息"],
    [
        ("提交文件名", "翼揽无余.apk"),
        ("文件大小", "1,449,345,968 字节（约 1.35 GiB）"),
        ("应用包名", "com.TianWenZhiDa.pico4"),
        ("版本", "Version Name 0.1 / Version Code 1"),
        ("处理器架构", "ARM64-v8a"),
        ("Android要求", "最低 API 29；目标 API 35"),
        ("图形能力", "OpenGL ES 3.0"),
        ("SHA-256", "E3BDACC4F88C7C300F72D4203F175E330FE75ECA063E89330B4B091F920317E8"),
    ],
    [2700, 6660],
    first_col_bold=True,
)

add_callout(
    doc,
    "提交前提醒",
    "当前APK安装后的应用显示名称为“pico 4”，不是“翼揽无余”。建议最终提交前在 Unity 的“项目设置 > Player > Product Name”中改为“翼揽无余”，并将版本号更新为 1.0 后重新构建。若不重新构建，评委需在应用库中查找名称“pico 4”。",
    fill=LIGHT_GOLD,
    accent=GOLD,
)

doc.add_heading("二、适用设备与运行环境", level=1)
add_table(
    doc,
    ["设备", "兼容状态", "说明"],
    [
        ("PICO 4 Ultra", "推荐/指定", "本项目目标设备；支持当前SDK、手势交互和视频透视功能。"),
        ("PICO 4 Ultra Enterprise", "建议实机复测", "硬件系列匹配，但提交前仍建议完成一次完整流程测试。"),
        ("PICO 4（非Ultra）", "不作为验收设备", "当前工程使用 PICO Integration 3.4.0 和视频透视；普通 PICO 4 的该能力存在SDK版本限制。"),
        ("PICO Neo3系列", "不支持本次验收", "没有按本项目的透视、手势和空间交互流程完成适配。"),
        ("Meta Quest/其他头显", "不支持", "当前APK为PICO平台构建，不能直接作为其他平台应用使用。"),
        ("安卓手机/平板", "不支持", "APK依赖头显追踪、PICO运行时和XR输入系统。"),
    ],
    [1900, 1900, 5560],
)

doc.add_heading("安装前准备", level=2)
add_bullet(doc, "一台 PICO 4 Ultra，建议升级至最新稳定系统；使用视频透视时系统版本应为 5.14.0 或以上。")
add_bullet(doc, "Windows 10 20H2 或以上电脑，并安装 PICO Developer Center 或 Android Platform Tools。")
add_bullet(doc, "一根支持数据传输的 USB-C 数据线；仅支持充电的数据线无法识别设备。")
add_bullet(doc, "头显建议至少预留 4 GB 可用存储空间，安装过程中不要断开数据线。")
add_bullet(doc, "体验区域应光线稳定、地面清晰并留出安全活动空间，便于手势追踪。")

doc.add_heading("三、开启开发者模式与USB调试", level=1)
add_numbered_sequence(doc, [
    "启动 PICO 4 Ultra，进入“设置 > 通用 > 关于本机”。",
    "将光标移至“软件版本号”，连续点击多次，直到左侧导航区域出现“开发者”选项。",
    "进入“开发者”，打开右上角的“USB调试”开关。",
    "使用USB数据线连接头显与电脑，在头显中接受USB调试授权；如出现“始终允许”选项，可按比赛电脑管理要求选择。",
    "打开 PICO Developer Center，确认设备状态显示为“已连接”。",
])

add_callout(
    doc,
    "PICO 4 Ultra补充设置",
    "若 PICO Developer Center 不能正常识别串流服务，可进入“设置 > 通用”，将“电脑互联自动发现”关闭后再重新打开。不同PICO OS版本的文字位置可能略有变化。",
    fill=LIGHT_BLUE,
    accent=BLUE,
)

h = doc.add_heading("四、使用ADB安装APK（推荐）", level=1)
add_body(doc, "以下方式最适合比赛材料验收，可明确确认设备连接、安装结果和应用包名。")

doc.add_heading("步骤1：准备英文路径", level=2)
add_body(doc, "由于部分Windows命令行工具对中文路径处理不一致，建议将APK复制到简短的英文目录，例如：")
add_command(doc, r"C:\APK\WingLanWuYu.apk")

doc.add_heading("步骤2：检查设备连接", level=2)
add_command(doc, "adb devices")
add_body(doc, "设备序列号后显示“device”表示连接正常；显示“unauthorized”时，应佩戴头显并确认USB调试授权。")

doc.add_heading("步骤3：执行安装", level=2)
add_command(doc, 'adb install -r "C:\\APK\\WingLanWuYu.apk"')
add_body(doc, "参数“-r”表示覆盖安装并尽量保留原有应用数据。由于APK约1.35 GiB，传输和安装可能需要较长时间；出现“Success”才表示安装完成。")

doc.add_heading("步骤4：处理签名冲突", level=2)
add_body(doc, "如果出现 INSTALL_FAILED_UPDATE_INCOMPATIBLE，说明设备上已有相同包名但签名不同的版本。先卸载旧版，再重新安装：")
add_command(doc, "adb uninstall com.TianWenZhiDa.pico4")
add_command(doc, 'adb install "C:\\APK\\WingLanWuYu.apk"')
add_callout(doc, "注意", "卸载旧版会删除该应用在设备上的本地数据。比赛验收设备通常可直接执行，但请先确认没有需要保留的测试记录。", fill=LIGHT_RED, accent=RED)

doc.add_heading("步骤5：验证安装", level=2)
add_command(doc, "adb shell pm list packages | findstr TianWenZhiDa")
add_body(doc, "若返回 package:com.TianWenZhiDa.pico4，表示应用包已经存在。也可以在头显应用库中查找当前显示名称“pico 4”。")

doc.add_heading("五、备用安装方式", level=1)
add_body(doc, "如果验收电脑不方便使用ADB，可尝试通过USB把APK复制到头显存储，再使用头显文件管理器打开并安装。首次侧载应用时，系统可能要求允许“安装未知来源应用”。由于不同PICO系统版本和企业管理策略可能限制该功能，比赛现场仍建议优先准备ADB安装方式。")

doc.add_heading("六、启动与首次运行", level=1)
add_numbered_sequence(doc, [
    "断开USB数据线或保持连接均可，在PICO应用库中找到“pico 4”（当前构建名称）并启动。",
    "首次启动会先显示Unity启动画面。项目资源较大，请等待片头场景出现，不要连续退出或重复启动。",
    "检查片头粒子、图片/视频和声音是否正常播放；片头结束后使用手势射线点击进入按钮。",
    "进入机型选择页面后，依次验证机型选择、核心交互、部件页面、爆炸图、模拟驾驶舱、虚拟课堂和答题页面。",
])

add_callout(doc, "网络说明", "当前比赛构建的核心模型、图片、视频和音频资源以本地内容为主，正常体验不依赖外网。若后续加入在线智能问答功能，则需要稳定的网络连接。", fill=LIGHT_GREEN, accent=GREEN)

doc.add_heading("七、基本操作说明", level=1)
add_table(
    doc,
    ["操作", "使用方式"],
    [
        ("手势射线", "将手掌保持在头显可识别范围内，用食指指向按钮或三维热点。"),
        ("确认点击", "食指与拇指捏合，完成按钮、标签或热点的选择。"),
        ("飞机旋转", "在支持双手旋转的场景中，双手同时捏合并改变两手连线方向。"),
        ("零件抓取", "在爆炸图中对准零件后捏合抓取，可移动、旋转或按场景设置进行缩放。"),
        ("驾驶舱操作", "使用射线或拖动方式操作摇杆和左右踏板，观察姿态、UI和语音反馈。"),
        ("页面返回", "使用页面左上角或场景中的返回按钮回到上一层。"),
    ],
    [2500, 6860],
    first_col_bold=True,
)

doc.add_heading("推荐验收顺序", level=2)
add_body(doc, "01片头 → 02机型选择 → 03核心交互 → 04部件认知 → 05爆炸拆解 → 06模拟驾驶舱 → 07虚拟课堂 → 08知识测试")

h = doc.add_heading("八、常见问题排查", level=1)
h.paragraph_format.page_break_before = True
add_table(
    doc,
    ["现象", "可能原因", "处理方式"],
    [
        ("adb devices没有设备", "数据线不支持传输或USB调试未开启", "更换数据线；重新开启USB调试；重插USB接口。"),
        ("显示unauthorized", "头显尚未授权电脑", "佩戴头显，在USB调试提示中点击允许。"),
        ("安装空间不足", "APK较大且安装需要临时空间", "删除不需要的应用或视频，建议保留至少4 GB空间。"),
        ("更新安装失败", "旧APK签名与当前APK不同", "卸载包名 com.TianWenZhiDa.pico4 后重新安装。"),
        ("应用库找不到翼揽无余", "当前安装显示名称仍是pico 4", "在应用库或未知来源区域查找“pico 4”。"),
        ("启动后暂时黑屏", "首次加载资源较慢或应用状态异常", "先等待加载；仍无内容时退出应用、重启头显并重新安装。"),
        ("没有手势射线", "手势追踪关闭、控制器仍激活或环境过暗", "开启手势追踪；放下控制器；改善环境光线并把双手移入视野。"),
        ("按钮变色但无法点击", "捏合动作未识别或交互状态未完成", "重新张开手指后再捏合，保持目标处于射线末端。"),
        ("没有声音", "系统音量低、静音或输出设备异常", "提高头显媒体音量，断开不需要的蓝牙音频设备并重启应用。"),
        ("透视/MR画面异常", "设备型号或系统版本不匹配", "使用PICO 4 Ultra，并升级到5.14.0或以上系统。"),
    ],
    [2300, 2800, 4260],
)

doc.add_heading("九、卸载方法", level=1)
add_command(doc, "adb uninstall com.TianWenZhiDa.pico4")
add_body(doc, "也可以在PICO应用库中进入应用详情后选择卸载。")

h = doc.add_heading("十、比赛提交前检查清单", level=1)
h.paragraph_format.page_break_before = True
checks = [
    "将 Unity Product Name 从“pico 4”修改为“翼揽无余”。",
    "将版本号从0.1调整为正式提交版本，例如1.0，并保留版本记录。",
    "确认最终APK文件名、文档中的大小和SHA-256与实际提交文件一致。",
    "在一台已卸载旧版本的PICO 4 Ultra上完成全新安装测试。",
    "从片头开始完整体验一次所有主要场景，确认不存在黑屏或无法返回的问题。",
    "检查左右手射线、捏合点击、双手旋转和爆炸零件交互。",
    "检查所有视频、旁白、背景音乐和操作提示音频。",
    "确认头显至少保留足够存储空间，并准备可传输数据的USB-C线。",
    "比赛提交包中同时放入APK和本安装指南。",
]
table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.LEFT
table.rows[0].cells[0].text = "状态"
table.rows[0].cells[1].text = "检查内容"
for cell in table.rows[0].cells:
    set_cell_shading(cell, NAVY)
    format_cell_text(cell, color=WHITE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
for item in checks:
    cells = table.add_row().cells
    cells[0].text = "□"
    cells[1].text = item
    format_cell_text(cells[0], size=12, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell_text(cells[1])
set_table_geometry(table, [900, 8460])
set_table_borders(table)

doc.add_heading("推荐提交目录", level=2)
add_command(doc, "翼揽无余_比赛提交包\\\n  翼揽无余.apk\n  翼揽无余_APK安装与运行指南.docx")

doc.add_heading("十一、参考资料", level=1)
add_body(doc, "PICO开发者中心快速开始：https://developer.picoxr.com/zh/document/unity/pdc-basic-info/")
add_body(doc, "PICO Interaction Sample安装说明：https://developer.picoxr.com/document/unity/pico-interaction-sample/")
add_body(doc, "PICO Video Seethrough设备要求：https://developer.picoxr.com/en/document/unity/seethrough/")

h = doc.add_heading("十二、提交信息（参赛团队填写）", level=1)
h.paragraph_format.page_break_before = True
add_table(
    doc,
    ["项目", "填写内容"],
    [
        ("作品名称", "翼揽无余"),
        ("参赛团队/学校", "____________________________"),
        ("负责人", "____________________________"),
        ("联系电话", "____________________________"),
        ("最终APK版本", "____________________________"),
        ("最终测试设备", "PICO 4 Ultra / 系统版本：____________"),
    ],
    [2700, 6660],
    first_col_bold=True,
)

# Keep tables from splitting rows and set metadata.
for table in doc.tables:
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)

doc.core_properties.title = "《翼揽无余》APK安装与运行指南"
doc.core_properties.subject = "PICO 4 Ultra 比赛作品安装说明"
doc.core_properties.author = "翼揽无余项目组"
doc.core_properties.keywords = "翼揽无余, PICO 4 Ultra, APK, 安装指南, Unity MR"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
